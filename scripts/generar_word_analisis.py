import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips (1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Colors
    COLOR_PRIMARY = RGBColor(24, 43, 73)      # Navy Deep
    COLOR_SECONDARY = RGBColor(41, 74, 110)  # Slate Navy
    COLOR_TEXT = RGBColor(30, 41, 59)        # Slate Dark
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate Muted
    COLOR_LISTO = RGBColor(22, 101, 52)      # Green Dark
    COLOR_PARCIAL = RGBColor(180, 83, 9)     # Amber Dark
    COLOR_FALTANTE = RGBColor(185, 28, 28)   # Crimson Dark

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # Document Header / Title
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_after = Pt(4)
    run_pre = p_pre.add_run("ERP ALMAHUE · INFORME DE ANÁLISIS Y ESTADO DE REQUERIMIENTOS")
    run_pre.font.size = Pt(9.5)
    run_pre.font.bold = True
    run_pre.font.color.rgb = COLOR_MUTED

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.space_before = Pt(0)
    run_title = p_title.add_run("Análisis del Sistema vs Grabaciones Presenciales del 26/10/2026")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Contraste de requerimientos solicitados en sala (sesión mañana 12.33 y sesión tarde 15.20) contra la implementación actual en el sistema (ERP Backend y Frontend).")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Box
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(16)
    r_meta1 = p_meta.add_run("Fecha de emisión: ")
    r_meta1.bold = True
    p_meta.add_run("27 de Octubre de 2026  |  ")
    r_meta2 = p_meta.add_run("Cliente / Empresa: ")
    r_meta2.bold = True
    p_meta.add_run("Almahue Export SpA  |  ")
    r_meta3 = p_meta.add_run("Participantes principales: ")
    r_meta3.bold = True
    p_meta.add_run("María Jesús (MJ), Lupe, Francisca, Consultores ERP")

    # Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("―" * 55)
    run_div.font.color.rgb = RGBColor(203, 213, 225)

    # 1. Introducción y Metodología
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Introducción y Metodología de Contraste")
    r_h1.font.size = Pt(15)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p_intro = doc.add_paragraph()
    p_intro.add_run(
        "El presente informe analiza los requerimientos y decisiones de diseño levantados durante las dos sesiones presenciales de trabajo en cliente realizadas el 26 de octubre de 2026:\n"
        "• Sesión de la Mañana (audio 26-10-2026 12.33.m4a, duración 01:24:40): Foco en Plan de Cuentas N:N, Reportería (Resultado del Ejercicio vs Situación Financiera al 31/12), Carga Masiva Agrosoft, Orden de Compra y Tesorería.\n"
        "• Sesión de la Tarde (audio 26-10-2026 15.20 REPARADO.m4a, duración ~01:54:00): Foco en Módulo Contratistas (parametrización actividad-labor, proformas, cierre contable mensual) y flujo de solicitud de apertura de mes y reversos.\n"
        "• Apuntes de Sala: Cuaderno de consultoría (revision en cliente.txt).\n\n"
        "Siguiendo las reglas canónicas de gobernanza del proyecto, las transcripciones de las declaraciones directas del cliente (María Jesús, Lupe) son la fuente primaria irrefutable sobre la cual se audita el código real del sistema."
    )

    # 2. Cuadro Resumen Ejecutivo
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)
    r_h2 = h2.add_run("2. Matriz Resumen de Requerimientos y Estado Actual")
    r_h2.font.size = Pt(15)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    summary_items = [
        ("1.1", "Dimensiones N:N por cuenta contable", "Parametrización", "MJ (12.33 [22:15]): 3 columnas, filtro dinámico en OC y asientos", "Listo"),
        ("1.2", "Contador de cuentas en CC / Elemento / Área", "Parametrización", "MJ (12.33 [29:17]): Contador numérico en mantenedor + detalle al clic", "Listo"),
        ("1.3", "Flags Agrosoft en Plan (PLT, VAR, VCTO, CF)", "Parametrización", "MJ (15.20 [03:39]): Quitar planta; variedad opcional en maestro", "Listo"),
        ("1.4", "Bancos con múltiples cuentas y moneda", "Parametrización", "Operación Almahue / Apunte P-3: Cuenta corriente CLP y USD por banco", "Listo"),
        ("1.5", "Selector multi-empresa al crear usuarios", "Administración", "Apunte A-1: Soporte para usuarios asignados a dos empresas del grupo", "Listo"),
        ("1.6", "Historial de parametrizaciones (antes/después)", "Parametrización", "MJ (15.20 [48:40]): Auditoría legible sin metadatos técnicos", "Parcialmente listo"),
        ("2.1", "Orden de columnas en OC y bloqueo sin cuenta", "Compras", "MJ (12.33 [44:04]): Cuenta > CC > Elemento; bloqueo si no hay cuenta", "Listo"),
        ("2.2", "Solicitud al administrador si dimensión no existe", "Compras", "MJ (12.33 [44:13]): Impedir alta libre; flujo formal con motivo", "Listo"),
        ("2.3", "Múltiples imputaciones por ítem en OC", "Compras", "MJ (12.33 [47:25]): Preferencia por múltiples líneas de imputación", "Listo"),
        ("2.4", "Detalle de OC en pantalla tipo árbol contable", "Compras", "MJ (12.33 [55:17]): Vista desglosada estilo flujo de caja", "Parcialmente listo"),
        ("2.5", "PDF de OC proveedor sin dimensiones contables", "Compras", "MJ (12.33 [50:24]): Hoja oficial proveedor sin CC/Elemento interno", "Faltante"),
        ("3.1", "Tipos de comprobante formalizados (6 tipos)", "Contabilidad", "Apunte C-6 y MJ (12.33 [10:03]): Apertura, Egresos, Ingresos, etc.", "Listo"),
        ("3.2", "Carga masiva comprobantes tipo Agrosoft", "Contabilidad", "MJ (12.33 [09:42]): Excel nativo con rechazo atómico fail-closed", "Parcialmente listo"),
        ("3.3", "Resultado del ejercicio y acumulado al 31/12", "Reportería", "MJ (12.33 [00:19]): Coexistencia obligatoria en Situación Financiera", "Faltante"),
        ("3.4", "Filtros Mayor/Diario ('Todos' = año en curso)", "Reportería", "Apunte C-12 / MJ: Selector ancho y restricción a ejercicio actual", "Parcialmente listo"),
        ("3.5", "Comprobante contabilizado directo y TC del día", "Contabilidad", "Apuntes C-8/C-10: Guardar contabilizado si está cuadrado", "Listo"),
        ("4.1", "Inversión jerarquía Actividad > Labores", "Contratistas", "MJ (15.20 [05:29]): Actividad agrupa labores; asociar desde actividad", "Parcialmente listo"),
        ("4.2", "Cadena de aprobación trasladada a Proforma", "Contratistas", "MJ (15.20 [23:28]): Retirar de captura diaria; aprobar proforma", "Parcialmente listo"),
        ("4.3", "Bloques de cierre mensual (Verde vs Gris)", "Contratistas", "MJ (15.20 [24:54]): Borde verde centraliza; gris informativo", "Parcialmente listo"),
        ("4.4", "Centralización contable y cuentas puente", "Contratistas", "MJ (15.20 [01:26]): Mano de obra vs Facturas por recibir", "Listo"),
        ("4.5", "Solicitud apertura de mes y reverso", "Contratistas", "MJ (15.20 [01:07]): Notificación con enlace directo a reversar asiento", "Parcialmente listo"),
        ("4.6", "Regla de cascada en apertura/reverso", "Contratistas", "MJ (15.20 [01:30]): Control entre periodos cerrados consecutivos", "Faltante"),
        ("4.7", "Edición acotada sin reverso si no hay factura", "Contratistas", "MJ (15.20 [01:36]): Modificación antes/después para cuartel/labor", "Listo"),
        ("5.1", "Parser cartola descarga directa Banco Chile", "Tesorería", "Lupe / MJ (12.33 [01:05]): Procesar archivo oficial web del banco", "Parcialmente listo"),
        ("6.1", "Honorarios, Nota de Cobro y OPB Embarcador", "Compras/Cont.", "Apuntes P-5 a P-7: No tratados en sala (12.33 [01:08:36])", "Faltante"),
    ]

    table = doc.add_table(rows=len(summary_items) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths
    col_widths = [Inches(0.5), Inches(2.2), Inches(1.2), Inches(2.0), Inches(1.1)]

    # Header Row
    headers = ["#", "Requerimiento Solicitado", "Módulo", "Contexto / Solicitante", "Estado"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "182B49")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    # Data Rows
    for r_idx, item in enumerate(summary_items, start=1):
        row_cells = table.rows[r_idx].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx in range(5):
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # #
        row_cells[0].text = item[0]
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.runs[0].font.size = Pt(9)
        p0.runs[0].font.bold = True

        # Requerimiento
        row_cells[1].text = item[1]
        p1 = row_cells[1].paragraphs[0]
        p1.runs[0].font.size = Pt(9.5)
        p1.runs[0].font.bold = True

        # Modulo
        row_cells[2].text = item[2]
        p2 = row_cells[2].paragraphs[0]
        p2.runs[0].font.size = Pt(9)

        # Contexto
        row_cells[3].text = item[3]
        p3 = row_cells[3].paragraphs[0]
        p3.runs[0].font.size = Pt(8.5)
        p3.runs[0].font.color.rgb = COLOR_MUTED

        # Estado
        estado = item[4]
        row_cells[4].text = estado
        p4 = row_cells[4].paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_st = p4.runs[0]
        run_st.font.size = Pt(9)
        run_st.font.bold = True
        if estado == "Listo":
            run_st.font.color.rgb = COLOR_LISTO
            set_cell_shading(row_cells[4], "DCFCE7") # Light Green
        elif estado == "Parcialmente listo":
            run_st.font.color.rgb = COLOR_PARCIAL
            set_cell_shading(row_cells[4], "FEF3C7") # Light Amber
        else:
            run_st.font.color.rgb = COLOR_FALTANTE
            set_cell_shading(row_cells[4], "FEE2E2") # Light Red

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Page Break after Summary Table
    doc.add_page_break()

    # Detailed Sections
    sections_data = [
        {
            "num": "3",
            "title": "Módulo 1: Parametrización y Plan de Cuentas",
            "items": [
                {
                    "title": "1.1 Matriz de Plan de Cuentas y Relaciones N:N por Cuenta",
                    "context": "En la sesión de la mañana ([22:15]–[25:59]), María Jesús (MJ) explicó que en Agrosoft una cuenta contable (ej. Gastos de Aduana y Embarque) no tiene una única dimensión, sino que puede aplicar a múltiples especies/centros (cerezas, manzanas, etc.) y múltiples elementos según el departamento (DUS, certificados, etc.). MJ aprobó la propuesta de una interfaz de tres columnas con agregar/quitar donde cada cuenta imputable vincule N centros de costo, N elementos y N áreas de negocio, y que al usar la cuenta en órdenes de compra o asientos contables el sistema filtre y ofrezca únicamente las dimensiones asociadas.",
                    "status": "Listo",
                    "evidence": "Implementado en modelos de unión CuentaCentroCosto, CuentaElementoCosto y CuentaAreaNegocio en backend (Prisma) con migración 20261026230000_banco_cuentas_solicitud_vinculo. En frontend, PlanCuentasPage.tsx cuenta con la pestaña y vista Matriz con scroll vertical y horizontal para asignar y remover dimensiones por cuenta sin perder la referencia jerárquica."
                },
                {
                    "title": "1.2 Contador de Cuentas Asociadas en Mantenedores de Dimensiones",
                    "context": "En la mañana ([29:17]–[29:59]), durante la revisión de los catálogos de Centro de Costo, Elemento de Costo y Área de Negocio, se consultó si convenía ver la relación inversa. MJ indicó que ver todas las cuentas desde el centro de costo sería engorroso, pero solicitó que en las tablas de mantenedores existiera una columna con un contador de cantidad de cuentas asociadas y que al hacer clic se pudiera revisar el detalle.",
                    "status": "Listo",
                    "evidence": "En catalogos.service.ts se incluye _count: { select: { cuentasVinculo: true } }, retornando cuentasAsociadas. En CatalogosPages.tsx y AreasNegocioPage.tsx se renderiza la columna 'Cuentas plan' con el contador numérico interactivo."
                },
                {
                    "title": "1.3 Flags de Requisitos Agrosoft en Cuentas (PLT, VAR, VCTO, CF)",
                    "context": "En la tarde ([02:47]–[04:10]), se revisaron las siglas de Agrosoft en la parametrización de cuentas contables. MJ aclaró que Planta (PLT) no se utiliza y debe eliminarse de la operación diaria; Variedad (VAR) tampoco se usa activamente, pero acordaron dejarla disponible en el maestro por si a futuro se requiere; mientras que Vencimiento (VCTO) y Código Financiero (CF) son requeridos por las cuentas de pasivo/proveedores.",
                    "status": "Listo",
                    "evidence": "Campos agregados en el modelo CuentaContable (requierePlanta, requiereVariedad, requiereVencimiento, requiereCodigoFinanciero) mediante la migración 20261027121500_cuenta_flags_plt_var_vcto_cf. Sincronización realizada mediante scripts/poblar_flags_agrosoft.js."
                },
                {
                    "title": "1.4 Banco con Múltiples Cuentas Bancarias y Moneda por Cuenta",
                    "context": "Apunte P-3/P-3b de la reunión. Almahue opera con bancos (principalmente Banco de Chile) que poseen más de una cuenta bancaria (ej. cuenta corriente en CLP y cuenta en USD) bajo la misma entidad bancaria. Se solicitó registrar número de cuenta, tipo (por defecto Cuenta Corriente) y moneda de cada cuenta.",
                    "status": "Listo",
                    "evidence": "Modelo BancoCuenta vinculado a Banco con campos numero, tipo (CUENTA_CORRIENTE), monedaCodigo y activo en backend. Administrador de cuentas bancarias por entidad bancaria operativo en BancosPage.tsx."
                },
                {
                    "title": "1.5 Selector de Empresa al Crear Usuarios",
                    "context": "Apunte A-1. Se observó que al crear usuarios en la administración solo se visualizaba una empresa en el formulario, a pesar de que los operadores pertenecen a dos empresas del grupo Almahue.",
                    "status": "Listo",
                    "evidence": "Corregido en EmpresasPage.tsx y selector multitenant EmpresaSelector para respetar las empresas asignadas al usuario."
                },
                {
                    "title": "1.6 Historial de Cambios en Parametrizaciones Generales",
                    "context": "En la tarde ([48:40]–[50:18]) y apunte P-4. Se reportó que el historial de parametrizaciones no era comprensible ni funcional para el usuario. MJ solicitó que registre acción, fecha, usuario responsable y un comparativo antes / después en lenguaje humano, dejando metadatos técnicos de base de datos restringidos a soporte.",
                    "status": "Parcialmente listo",
                    "evidence": "En el módulo de Contratistas está implementado en AuditoriaContratistasPage.tsx con formateador limpio auditoria-format.ts. En el resto de catálogos generales (Plan de Cuentas, Bancos, Proveedores) la vista consolidada de auditoría antes/después aún no está desplegada en la UI para el usuario final."
                }
            ]
        },
        {
            "num": "4",
            "title": "Módulo 2: Compras y Orden de Compra (OC)",
            "items": [
                {
                    "title": "2.1 Orden Estricto de Columnas en Ítems de OC y Bloqueo sin Cuenta",
                    "context": "En la mañana ([44:04]–[44:13]), MJ solicitó que en las líneas de la Orden de Compra el orden sea estrictamente: Cuenta Contable > Centro de Costo > Elemento de Costo. Además, si el usuario no ha seleccionado primero una cuenta contable, los selectores de centro de costo y elemento deben permanecer bloqueados/deshabilitados, impidiendo imputaciones erróneas.",
                    "status": "Listo",
                    "evidence": "En OrdenCompraWizardPage.tsx, las columnas de cada ítem se ordenan: Descripción → Cuenta contable → Centro de costo → Elemento de costo. El selector de Elemento de Costo tiene disabled={!it.cuentaContableId}, y el Centro de Costo filtra dinámicamente mediante catalogoDimensionOc según las cuentas vinculadas."
                },
                {
                    "title": "2.2 Solicitud al Administrador si la Dimensión no está Vinculada",
                    "context": "En la mañana ([44:13]–[45:24]), se debatió qué ocurre si un operador al armar una OC necesita un centro de costo o elemento que no está asociado a la cuenta en el plan. MJ fue enfática en que el operador no debe poder crearlo libremente, ya que desordenaría el presupuesto. Debe existir una opción en el selector que diga 'Solicitar al administrador...', abriendo un flujo donde el operador indica el motivo y la solicitud le llega al administrador/jefatura contable para su aprobación y vinculación.",
                    "status": "Listo",
                    "evidence": "Backend: Modelo SolicitudVinculoCuenta, endpoints POST /solicitudes-vinculo-cuenta, POST /solicitudes-vinculo-cuenta/:id/aceptar, etc., en solicitudes-vinculo-cuenta.service.ts. Frontend: En OrdenCompraWizardPage.tsx la opción lanza el modal de solicitud, y se gestiona en NotificacionesPage.tsx y MisSolicitudesPage.tsx."
                },
                {
                    "title": "2.3 Múltiples Imputaciones por Ítem (Líneas Separadas)",
                    "context": "En la mañana ([33:42]–[35:17] y [47:25]), se revisó el caso de compras que se dividen entre dos áreas (ej. software 50% TI y 50% Contabilidad). MJ aclaró que no le complica que vengan varias líneas en el detalle de la OC con la misma cuenta y distinto centro de costo, prefiriendo crear múltiples líneas en vez de submecanismos complejos dentro de un mismo ítem.",
                    "status": "Listo",
                    "evidence": "Nativo en el diseño de líneas de OrdenCompraWizardPage.tsx, permitiendo agregar N líneas con el mismo concepto o artículo e imputaciones independientes."
                },
                {
                    "title": "2.4 Vista de Detalle de OC en Pantalla (Estilo Flujo de Caja)",
                    "context": "En la mañana ([55:17]–[55:29]) y apunte C-4. MJ señaló que en Agrosoft la primera vista no muestra el centro de costo y tienen que abrir la 'segunda hoja' contable. Pidió que en el ERP la visualización del detalle de la OC en pantalla permita ver un desglose ordenado tipo árbol o tabla (cuenta y centros de costo), similar a la vista que se usa en el reporte de flujo de caja/estado de cuenta.",
                    "status": "Parcialmente listo",
                    "evidence": "El wizard y listado muestran las dimensiones por línea, pero el modal de detalle agrupado en árbol contable está pendiente de implementación visual dedicada."
                },
                {
                    "title": "2.5 Plantilla PDF de OC para el Proveedor sin Dimensiones Contables",
                    "context": "En la mañana ([50:24]–[53:30]), MJ explicó que en Agrosoft la orden de compra tiene dos hojas: la hoja 1 es la oficial que se envía al proveedor (artículo, descripción, cantidad, precio), la cual no debe mostrar centros de costo ni elementos internos ('no tienen sentido para el proveedor'). La hoja 2 es la contable interna. MJ pidió que el PDF descargable/enviable al proveedor contenga únicamente la hoja oficial sin las dimensiones internas.",
                    "status": "Faltante",
                    "evidence": "El ERP actual cuenta con gestión de adjuntos y trazabilidad, pero no posee aún el generador de PDF de la OC con plantilla dual o formato proveedor limpio sin dimensiones contables."
                }
            ]
        },
        {
            "num": "5",
            "title": "Módulo 3: Contabilidad, Asientos y Reportería",
            "items": [
                {
                    "title": "3.1 Tipos de Comprobante Formalizados en Creación y Módulos",
                    "context": "Apuntes C-6/C-7 y sesión de la mañana ([10:03]). En Agrosoft los comprobantes se clasifican formalmente según su origen. Se acordaron seis tipos estándar: Apertura, Egresos, Ingresos, Proveedores, Traspasos, Ventas. Si se genera manualmente desde contabilidad, el usuario selecciona el tipo; si proviene de un módulo operativo (ej. cierre de contratistas o libro de compras), se asigna automáticamente.",
                    "status": "Listo",
                    "evidence": "Definido en tipo-comprobante.ts con la constante TIPOS_COMPROBANTE_NUEVO y cableado en el modal de nuevo comprobante en AsientosPage.tsx."
                },
                {
                    "title": "3.2 Carga Masiva de Comprobantes tipo Agrosoft (Validación Fail-Closed)",
                    "context": "En la mañana ([09:42]–[16:57]), MJ mostró el archivo Excel real de carga masiva de Agrosoft. Consta de columnas fijas (correlativo de líneas, tipo de comprobante, fecha, TC) y columnas que se exigen según la cuenta contable (CC, área, auxiliar, referencia, etc.). MJ insistió en un formato restrictivo y fail-closed: si una cuenta exige centro de costo y la fila no lo tiene, o si la suma de Debe ≠ Haber o la fecha no calza con el periodo, el sistema no debe cargar nada del lote; debe arrojar error inmediato para corregir el archivo y volver a subirlo íntegro.",
                    "status": "Parcialmente listo",
                    "evidence": "Existe el endpoint POST /asientos/bulk en contabilidad.service.ts, pero procesa un DTO JSON fila por fila. Falta el parser directo del Excel nativo de Agrosoft con validación atómica estricta (rechazo integral del lote si una fila falla o no cuadra)."
                },
                {
                    "title": "3.3 Tratamiento de Resultado del Ejercicio y Situación Financiera al 31/12",
                    "context": "En la mañana ([00:19]–[02:03]), MJ explicó con balances reales en pantalla la regla contable clave: el Estado de Resultados de un mes (ej. enero) muestra solo las cuentas de resultado de ese mes. El Resultado del Ejercicio del año anterior pasa el 31 de diciembre al Estado de Situación Financiera como Resultado Acumulado. En el Balance de Situación Financiera deben coexistir ambas líneas (Resultado del Ejercicio en curso y Resultado Acumulado de años anteriores); de lo contrario, el balance no cuadra.",
                    "status": "Faltante",
                    "evidence": "En EstadosFinancierosPage.tsx, el cálculo del balance inyecta únicamente el resultado del periodo consultado (re), pero no arrastra ni coexiste con el Resultado Acumulado al cierre del 31/12 para ejercicios sucesivos."
                },
                {
                    "title": "3.4 Filtros en Libro Mayor, Diario y Balance (Año en Curso vs Histórico)",
                    "context": "Apuntes C-12 a C-17 y discusión de reportería. Se solicitó que los reportes permitan filtrar por mes y año, y que la opción 'Todos' en la vista rápida filtre únicamente los periodos del año en curso (no toda la historia de la BD). Además, el selector de cuentas en el Mayor debe tener ancho suficiente para no truncar nombres.",
                    "status": "Parcialmente listo",
                    "evidence": "En LibroDiarioMayorPages.tsx, cuando periodoVista.todo está activo, consulta todos los periodos existentes en la base de datos sin restringir por defecto al año fiscal en curso."
                },
                {
                    "title": "3.5 Creación de Comprobante Contabilizado Directo y Validación TC",
                    "context": "Apuntes C-8, C-9 y C-10. Al crear un comprobante manual cuadrado en contabilidad, debe guardarse directamente como CONTABILIZADO. Si falta el tipo de cambio del día, debe permitir ingreso manual y levantar una alerta para el administrador.",
                    "status": "Listo",
                    "evidence": "Los comprobantes manuales válidos se guardan directamente como contabilizados con generación de correlativo formal; la revaluación y consulta de TC opera vía revaluacion-tc.util.ts."
                }
            ]
        },
        {
            "num": "6",
            "title": "Módulo 4: Contratistas (Sesión de la Tarde)",
            "items": [
                {
                    "title": "4.1 Inversión de Jerarquía en Parametrización: Actividad > Labores",
                    "context": "En la tarde ([05:29]–[09:13]), al revisar la parametrización de labores y actividades, MJ observó que la interfaz pedía seleccionar una labor para asociarle actividades. MJ aclaró categóricamente: 'Va al revés. La actividad manda, una actividad agrupa 20 labores... la actividad no se repite y la labor pertenece a esa actividad'. Pidió que la asociación se realice seleccionando la actividad y vinculando a ella las labores correspondientes.",
                    "status": "Parcialmente listo",
                    "evidence": "En ParametrizacionContratistasPage.tsx, la pestaña asociaciones todavía mantiene el estado laborId y busca actividades para vincular, en lugar de manejar la pantalla desde actividadId hacia la selección de sus labores exclusivas."
                },
                {
                    "title": "4.2 Cadena de Aprobación Trasladada a la Proforma",
                    "context": "En la tarde ([23:28]–[24:11]), Carlos y Sergio mostraron que el cálculo de la cadena de aprobación se estaba ejecutando en el formulario de Ingreso Labor Diario. MJ intervino inmediatamente: 'El ingreso diario no, sería muy lento para los chiquillos... la proforma es la que necesita aprobación'. El ingreso diario es una captura operativa masiva de faena; el documento que representa un monto consolidado y requiere firma de jefatura es la Proforma.",
                    "status": "Parcialmente listo",
                    "evidence": "En backend la proforma tiene su flujo de aprobación consolidado. Sin embargo, en el frontend, el componente CadenaLaborPreview sigue presente en el modal de ingreso diario y debe ser retirado para que la cadena se visualice exclusivamente en la generación de la Proforma."
                },
                {
                    "title": "4.3 Bloques Visuales de Cierre Mensual (Verde = Centralizado, Gris = Informativo)",
                    "context": "En la tarde ([24:54]–[25:23]), MJ y el equipo revisaron la vista previa del cierre de mes de contratistas. Se acordó distinguir nítidamente: un bloque principal con borde verde que contiene las proformas aprobadas que entrarán al asiento contable de centralización, y dos bloques informativos grises que no entran al asiento (proformas en borrador sin aprobación e ingresos diarios sin proforma asociada), para alertar si alguien olvidó asociar faenas.",
                    "status": "Parcialmente listo",
                    "evidence": "En TraspasoContratistasResumenPanel.tsx existen las tarjetas KPI ('Entran al asiento', 'Fuera · sin aprobación', 'Fuera · ingresos'), pero aún falta la separación visual de bloques verde vs gris solicitada en sala."
                },
                {
                    "title": "4.4 Centralización Contable y Enlace con Compras (Cuentas Puente)",
                    "context": "En la tarde ([01:26]–[01:29]), MJ explicó con exactitud la contabilidad de contratistas: al aprobar la proforma se reconoce el gasto del mes (Mano de Obra Contratista al Debe contra Facturas por Recibir Contratistas al Haber). Al centralizar a fin de mes, se genera un gran asiento en contabilidad con todas las proformas aprobadas. Luego, cuando el contratista entrega la factura (que puede llegar al mes siguiente), en el Libro de Compras esa factura netea la cuenta puente Facturas por Recibir Contratistas contra Proveedores.",
                    "status": "Listo",
                    "evidence": "El servicio de centralización genera el asiento agrupado contra las cuentas puente parametrizadas. La asociación entre la proforma y la factura de compras está implementada en ProformasContratistaNivel1Page.tsx."
                },
                {
                    "title": "4.5 Flujo de Solicitud de Apertura de Mes y Reverso desde Módulo Origen",
                    "context": "En la tarde ([01:07]–[01:17] y [01:33]–[01:35]). Un operador de abastecimiento/contratistas no tiene acceso a abrir periodos contables. Si necesita corregir una proforma con el mes cerrado, debe presionar 'Solicitar apertura' con un motivo; esta solicitud debe generar una notificación a Contabilidad (MJ); al hacer clic en la notificación, el sistema debe abrir directamente el asiento contable de centralización con un botón habilitado de Reversar que guarde la trazabilidad (quién solicitó y quién reversó).",
                    "status": "Parcialmente listo",
                    "evidence": "Existe el modelo PeriodoContableEvento que registra el historial de aperturas y cierres con motivo y usuario. Falta el botón de solicitud en el documento operativo que envíe la notificación in-app directa a Contabilidad con redirección al asiento específico para su reversa."
                },
                {
                    "title": "4.6 Regla de Cascada en Apertura de Meses y Reversos",
                    "context": "En la tarde ([01:30]–[01:31]). Se consultó si al equivocarse en un mes anterior (ej. mayo estando en septiembre) se puede abrir solo ese mes o debe ser en cascada. MJ explicó que en Agrosoft es una restricción en cascada, pero acordaron que en el ERP se debe evaluar si el reverso altera o no meses cerrados posteriores para evitar inconsistencias contables entre periodos consecutivos.",
                    "status": "Faltante",
                    "evidence": "Actualmente la apertura de periodos en contabilidad.service.ts no valida si existen periodos cerrados posteriores con arrastre de saldos afectados."
                },
                {
                    "title": "4.7 Edición Acotada sin Reverso Total (Solicitud de Modificación)",
                    "context": "En la tarde ([01:36]–[01:42]). MJ y el consultor analizaron que en la mayoría de los casos los errores en faenas son imputaciones no monetarias (ej. se equivocaron de cuartel o nombre de labor, pero no en el monto a pagar). Para evitar reversar toda la contabilidad y la factura, acordaron un flujo donde, siempre que la proforma NO tenga factura asociada, el usuario solicita una modificación indicando el antes y el después para aprobación de la jefatura de contratistas sin abrir todo el mes contable.",
                    "status": "Listo",
                    "evidence": "Implementado en el modal ProformaSolicitarModificacionModal.tsx que captura los cambios campo a campo (cantidad, precioUnitario, centroCostoId) con motivo y lo envía a revisión."
                }
            ]
        },
        {
            "num": "7",
            "title": "Módulo 5: Tesorería",
            "items": [
                {
                    "title": "5.1 Parser de Cartola Bancaria Descarga Directa Banco Chile",
                    "context": "En la mañana ([01:02]–[01:09]), Lupe y MJ discutieron sobre la carga de la cartola. Se confirmó que el formato debe ser el archivo oficial de descarga directa que entrega el portal de Banco de Chile (Excel/CSV sin edición manual previa por parte del usuario).",
                    "status": "Parcialmente listo",
                    "evidence": "En almahue-web.ts existe un parser heurístico robusto multi-hoja que detecta Fecha, Descripción, Cargos y Abonos. En banco-chile.ts, la función mapBancoChileRows indica 'Parser Banco Chile pendiente de muestra', esperando un archivo fixture real del portal web bancario para cerrar el parsing 100% nativo."
                }
            ]
        },
        {
            "num": "8",
            "title": "Módulo 6: Temas Fuera de Grabación / Pendientes de Definición",
            "items": [
                {
                    "title": "6.1 Boleta de Honorarios, Nota de Cobro y OPB Embarcador",
                    "context": "Los tres temas figuraban en los apuntes manuscritos (revision en cliente.txt), pero en la grabación de la mañana ([01:08:36]) se confirmó explícitamente: 'boleta de honorarios no lo vimos'. Se mantienen como pendientes de diseño operativo:\n1. Boleta de Honorarios: El ERP no emite boletas al SII; solo registra las recibidas y su retención.\n2. Nota de Cobro: Rendición de gastos internos que asocia facturas electrónicas desde Compras/Contabilidad.\n3. OPB (Embarcador): Intermediario que liquida pagos en el exterior y Almahue ingresa el detalle.",
                    "status": "Faltante",
                    "evidence": "Identificados formalmente en PLAN-IMPLEMENTACION-MODULOS-2026-10-26.md (puntos P-5 a P-7) y validacion-puntos-plan-2026-10-27.md como ítems pendientes de diagrama de flujo antes de codificar."
                }
            ]
        }
    ]

    for sec in sections_data:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(8)
        r_h = h.add_run(f"{sec['num']}. {sec['title']}")
        r_h.font.size = Pt(14)
        r_h.font.bold = True
        r_h.font.color.rgb = COLOR_PRIMARY

        for it in sec['items']:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.space_before = Pt(8)
            p_it.paragraph_format.space_after = Pt(2)
            r_it = p_it.add_run(it['title'])
            r_it.font.size = Pt(12)
            r_it.font.bold = True
            r_it.font.color.rgb = COLOR_SECONDARY

            # Status Badge Paragraph
            p_badge = doc.add_paragraph()
            p_badge.paragraph_format.space_before = Pt(0)
            p_badge.paragraph_format.space_after = Pt(4)
            r_st_label = p_badge.add_run("Estado en el Sistema: ")
            r_st_label.font.bold = True
            r_st_label.font.size = Pt(10)
            
            r_st_val = p_badge.add_run(f"[{it['status'].upper()}]")
            r_st_val.font.bold = True
            r_st_val.font.size = Pt(10)
            if it['status'] == "Listo":
                r_st_val.font.color.rgb = COLOR_LISTO
            elif it['status'] == "Parcialmente listo":
                r_st_val.font.color.rgb = COLOR_PARCIAL
            else:
                r_st_val.font.color.rgb = COLOR_FALTANTE

            # Context
            p_ctx = doc.add_paragraph()
            p_ctx.paragraph_format.space_before = Pt(0)
            p_ctx.paragraph_format.space_after = Pt(3)
            r_ctx_lbl = p_ctx.add_run("Contexto de la reunión: ")
            r_ctx_lbl.font.bold = True
            r_ctx_lbl.font.size = Pt(10)
            r_ctx_txt = p_ctx.add_run(it['context'])
            r_ctx_txt.font.size = Pt(10)

            # Evidence
            p_evi = doc.add_paragraph()
            p_evi.paragraph_format.space_before = Pt(0)
            p_evi.paragraph_format.space_after = Pt(10)
            r_evi_lbl = p_evi.add_run("Evidencia en el código: ")
            r_evi_lbl.font.bold = True
            r_evi_lbl.font.size = Pt(10)
            r_evi_txt = p_evi.add_run(it['evidence'])
            r_evi_txt.font.size = Pt(10)
            r_evi_txt.font.color.rgb = RGBColor(51, 65, 85)

    # 9. Conclusiones y Próximos Pasos
    h_concl = doc.add_paragraph()
    h_concl.paragraph_format.space_before = Pt(18)
    h_concl.paragraph_format.space_after = Pt(6)
    r_concl = h_concl.add_run("9. Conclusiones y Priorización de Trabajo")
    r_concl.font.size = Pt(14)
    r_concl.font.bold = True
    r_concl.font.color.rgb = COLOR_PRIMARY

    p_concl = doc.add_paragraph()
    p_concl.add_run(
        "1. Alto Avance Estructural en Parametrización y Compras: La arquitectura central solicitada por María Jesús (Plan de cuentas N:N con matriz interactiva, contador de vínculos, flags de Agrosoft y flujo de solicitud al administrador en OC) se encuentra completamente modelada y operativa.\n"
        "2. Focos Faltantes Críticos de Reportería: La coexistencia obligatoria entre 'Resultado del Ejercicio' y 'Resultado Acumulado' al 31/12 en el Balance de Situación Financiera es indispensable para el cuadre contable formal.\n"
        "3. Ajustes de Usabilidad en Contratistas: Invertir el flujo de asociación (Actividad agrupa Labores) y suprimir la visualización de aprobación en la captura diaria para concentrarla exclusivamente en la Proforma.\n"
        "4. Carga Masiva Fail-Closed: Incorporar el validador atómico para planillas nativas de Agrosoft que rechace el lote completo ante inconsistencias de Debe/Haber o ausencia de centros de costo obligatorios."
    )

    output_path = os.path.abspath(r"e:\source\repos\Almahue\docs\requerimientos\analisis-reuniones-2026-10-26.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
