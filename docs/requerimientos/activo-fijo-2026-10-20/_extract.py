import pathlib
import pymupdf
from openpyxl import load_workbook

base = pathlib.Path(__file__).resolve().parent

for name in ["MODULO DE ACTIVO FIJO.pdf", "RepMaestroAF.pdf"]:
    p = base / name
    doc = pymupdf.open(p)
    stem = p.stem.replace(" ", "-")
    parts = [f"# {name}", f"paginas: {doc.page_count}", ""]
    for i, page in enumerate(doc, 1):
        t = page.get_text("text")
        parts.append(f"## Pagina {i}")
        parts.append(t.strip() or "(sin texto / diagrama)")
        parts.append("")
        if i <= 12:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(1.35, 1.35))
            pix.save(str(base / f"{stem}-p{i:02d}.png"))
    (base / f"{stem}-texto.md").write_text("\n".join(parts), encoding="utf-8")
    print(name, "pages", doc.page_count)

xlsx = base / "09.-Deprciacion Activo Fijo SEPTIEMBRE.xlsx"
wb = load_workbook(xlsx, read_only=True, data_only=True)
lines = [f"# {xlsx.name}", f"hojas: {wb.sheetnames}", ""]
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    lines.append(f"## Hoja {sheet_name}")
    rows_out = 0
    for i, row in enumerate(ws.iter_rows(max_row=40, values_only=True), 1):
        vals = ["" if v is None else str(v).replace("\n", " ") for v in row]
        if any(v.strip() for v in vals):
            lines.append("| " + " | ".join(vals[:18]) + " |")
            rows_out += 1
    lines.append(f"(filas con dato en primeras 40: {rows_out})")
    lines.append("")
wb.close()
(base / "depreciacion-septiembre-muestra.md").write_text("\n".join(lines), encoding="utf-8")
print("excel ok")
